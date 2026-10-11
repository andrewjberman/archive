"""AMGT_CORPUS round 47, CRPS-aw-1 -- the team-sheet reader labels field-code rows in words (cited), classes override and
fund-year rows CONTROL by name, and carries each team cell's own Excel number format. Issuer's file. READ-ONLY (a workbook built
in memory; no S3, no network).

  PYTHONPATH=. python aw1_check.py          prints  `aw1 build <b>/4 control <c>/2`

The rules under test (CTO round-46 analysis ADDENDUM 1, msg 3972; ROUND_ANALYSIS.round46.md s.2-3):
  * a counted row whose label is `*_OVERRIDE` (regex `^[A-Z][A-Z0-9]*(_[A-Z0-9]+)*_OVERRIDE$`) or is `EQY_FUND_YEAR` is CONTROL
    wherever it sits and whatever its cells hold; it is never the period header and never a FIGURE or a HEADING;
  * `row_labels[i]` (aligned with `rows`) is None unless the label has the identifier form (`^[A-Z][A-Z0-9]*(_[A-Z0-9]+)+$`); for
    an identifier label it is {"text","source","claim","cite","code","why"}:
      - the code is a mnemonic in integrations/bloomberg/research/field_map.json: text = that field's `basis.definition` AS STORED,
        source "bloomberg", claim "held", cite = the field map's repo path, why None;
      - the code starts FF_ FFI_ FG_ FE_ or P_ (FactSet data items): text "FactSet line", source "factset", claim "absent", cite
        None, why states that no published FactSet description was found and the meaning is not inferred from the name;
      - any other identifier: text "Bloomberg line", source "bloomberg", claim "absent", cite None, why states the estate's field
        metadata does not describe it;
  * `row_formats[i]` is None unless row i is FIGURE; for a FIGURE row it is the row's cells' Excel number format strings in the
    period columns, aligned with `row_values[i]` ('General' where the workbook says General, None where the sheet carries no
    cell at all).
Everything round 46 serves is unchanged (rows, row_classes, row_numbers, row_values, periods, row_counts).
"""
import io
import json
import os
import sys

import openpyxl

from lib.research.corpus import dossier, team_sheet

HEADER_ROW = 9
YEARS = (2016, 2017, 2018, 2019, 2020)
PERIOD_COL0 = 5  # column E, 1-based
FIELD_MAP = "integrations/bloomberg/research/field_map.json"


def field_definition(code):
    top = os.popen("git rev-parse --show-toplevel").read().strip()
    with open(os.path.join(top, FIELD_MAP), encoding="utf-8") as fh:
        fields = json.load(fh)["fields"]
    for f in fields:
        if f["mnemonic"] == code:
            return f["basis"]["definition"]
    return None


def workbook_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model NEW VERSION"

    def put(row, col, value, fmt=None):
        c = ws.cell(row=row, column=col, value=value)
        if fmt:
            c.number_format = fmt

    put(2, 3, "Ticker:")
    put(2, 12, 1.392)
    put(2, 15, 497)
    put(3, 1, "FG_FACTSET_SECTOR")
    put(3, 9, 28.05)
    put(3, 15, 16118)
    put(4, 1, "EQY_FUND_YEAR")
    put(4, 9, 2020)
    put(4, 12, 2021)
    put(4, 15, 2022)
    put(HEADER_ROW, 3, "Financial Summary")
    for k, year in enumerate(YEARS):
        put(HEADER_ROW, PERIOD_COL0 + k, year)
    rows = [
        # label, values, number format of every period cell
        ("Revenue", [7976, 8671, 8700, 9009, 8753], "#,##0"),
        ("  y/y Growth", [None, 0.087, 0.003, 0.036, -0.028], "0.0%"),
        ("EPS", [1.2345, 1.3456, 1.4567, 1.5678, 1.6789], "0.00"),
        ("Capex", [310, 320, 330, 340, 350], None),
        ("BEST_EBIT", [5000.5, 5100.5, 5200.5, 5300.5, 5400.5], "#,##0"),
        ("FF_SALES", [9000, 9100, 9200, 9300, 9400], "#,##0"),
        ("BS_TOT_LOAN", [100, 110, 120, 130, 140], "General"),
        ("BEST_FPERIOD_OVERRIDE", [1, 2, 3, 4, 5], "General"),
    ]
    row = HEADER_ROW + 1
    for label, values, fmt in rows:
        put(row, 3, label)
        for k, v in enumerate(values):
            if v is not None:
                put(row, PERIOD_COL0 + k, v, fmt)
        row += 1
    # EQY_FUND_YEAR below the header, values the increasing years: an all-year row (round 46: HEADING), CONTROL by name now
    put(row, 3, "EQY_FUND_YEAR")
    for k, year in enumerate(YEARS):
        put(row, PERIOD_COL0 + k, year)
    other = wb.create_sheet("Model")
    for i in range(1, 13):
        other.cell(row=i, column=1, value="Old line %d" % i)
        other.cell(row=i, column=2, value=i)
        other.cell(row=i, column=3, value=i * 2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


WANT_ROWS = [
    "Ticker:", "FG_FACTSET_SECTOR", "EQY_FUND_YEAR", "Financial Summary", "Revenue", "y/y Growth", "EPS", "Capex",
    "BEST_EBIT", "FF_SALES", "BS_TOT_LOAN", "BEST_FPERIOD_OVERRIDE", "EQY_FUND_YEAR",
]
WANT_CLASSES = [
    "HEADING", "CONTROL", "CONTROL", "HEADING", "FIGURE", "FIGURE", "FIGURE", "FIGURE",
    "FIGURE", "FIGURE", "FIGURE", "CONTROL", "CONTROL",
]


def main():
    body = workbook_bytes()
    b1 = b2 = b3 = b4 = 0
    reading = None
    try:
        reading = team_sheet.sheet_reading(body).get("Model NEW VERSION")
    except Exception as exc:  # noqa: BLE001
        print("aw1 sheet_reading threw: %s" % str(exc)[:160])
    if reading:
        rows = reading.get("rows") or []
        classes = reading.get("row_classes") or []
        b1 = 1 if (rows == WANT_ROWS and classes == WANT_CLASSES) else 0
        print("aw1 b1 classes (override and fund-year rows CONTROL by name) equal expected %s" % (b1 == 1))
        if b1 != 1:
            print("aw1 b1 got %s" % (list(zip(rows, classes)),))
        labels = reading.get("row_labels")
        ok = isinstance(labels, list) and len(labels) == len(rows)
        if ok:
            by_row = {}
            for i, (lab, ent) in enumerate(zip(rows, labels)):
                by_row.setdefault(lab, ent)
            bloom = by_row.get("BEST_EBIT") or {}
            fact = by_row.get("FF_SALES") or {}
            lone = by_row.get("BS_TOT_LOAN") or {}
            want_text = field_definition("BEST_EBIT")
            ok = (
                want_text is not None
                and by_row.get("Revenue") is None and by_row.get("Ticker:") is None
                and bloom.get("text") == want_text and bloom.get("source") == "bloomberg" and bloom.get("claim") == "held"
                and str(bloom.get("cite") or "").endswith("field_map.json") and bloom.get("code") == "BEST_EBIT"
                and fact.get("text") == "FactSet line" and fact.get("source") == "factset" and fact.get("claim") == "absent"
                and fact.get("cite") is None and fact.get("code") == "FF_SALES"
                and "no published" in str(fact.get("why") or "").lower()
                and lone.get("text") == "Bloomberg line" and lone.get("source") == "bloomberg" and lone.get("claim") == "absent"
                and lone.get("cite") is None and lone.get("code") == "BS_TOT_LOAN" and str(lone.get("why") or "").strip() != ""
            )
        b2 = 1 if ok else 0
        print("aw1 b2 row_labels (Bloomberg description as stored, cited; FactSet line; Bloomberg line) %s" % (b2 == 1))
        formats = reading.get("row_formats")
        values = reading.get("row_values") or []
        ok = isinstance(formats, list) and len(formats) == len(rows)
        if ok:
            by = {lab: f for lab, f in zip(rows, formats)}
            ok = (
                by.get("Revenue") == ["#,##0"] * 5
                and by.get("y/y Growth") == [None, "0.0%", "0.0%", "0.0%", "0.0%"]
                and by.get("EPS") == ["0.00"] * 5
                and by.get("Capex") == ["General"] * 5
                and by.get("Ticker:") is None and by.get("BEST_FPERIOD_OVERRIDE") is None
                and all(len(f) == len(v) for f, v in zip(formats, values) if f is not None)
            )
        b3 = 1 if ok else 0
        print("aw1 b3 row_formats aligned with row_values, per period cell, 'General' where General, None where absent %s" % (b3 == 1))
        periods = reading.get("periods") or []
        counts = {c: sum(1 for k in classes if k == c) for c in ("FIGURE", "HEADING", "CONTROL")}
        rev = values[WANT_ROWS.index("Revenue")] if values else None
        b4 = 1 if (rev == [7976, 8671, 8700, 9009, 8753]
                   and [p.get("label") for p in periods] == [str(y) for y in YEARS]
                   and [p.get("column") for p in periods] == ["E", "F", "G", "H", "I"]
                   and counts == {"FIGURE": 7, "HEADING": 2, "CONTROL": 4}
                   and sum(counts.values()) == len(rows)) else 0
        print("aw1 b4 Revenue verbatim, periods and class counts %s sum %d rows %d" % (counts, sum(counts.values()), len(rows)))
    # control c1: `rows` is still the existing rule's counted labels in sheet order, and round 46's value keys are verbatim
    c1 = 0
    try:
        base_rows = team_sheet.sheet_labels(body).get("Model NEW VERSION")
        c1 = 1 if base_rows == WANT_ROWS else 0
        print("aw1 c1 sheet_labels unchanged %s (%d labels)" % (c1 == 1, len(base_rows or [])))
    except Exception as exc:  # noqa: BLE001
        print("aw1 c1 threw: %s" % str(exc)[:160])
    # control c2: both sheets still carried (model_rows_beside rule)
    c2 = 0
    try:
        both = team_sheet.sheet_labels(body)
        c2 = 1 if sorted(both) == ["Model", "Model NEW VERSION"] and len(both["Model"]) == 12 else 0
        print("aw1 c2 both sheets carried %s" % (c2 == 1))
    except Exception as exc:  # noqa: BLE001
        print("aw1 c2 threw: %s" % str(exc)[:160])
    print("aw1 METHOD_VERSION %s" % dossier.METHOD_VERSION)
    print("aw1 build %d/4 control %d/2" % (b1 + b2 + b3 + b4, c1 + c2))
    return 0 if (b1 + b2 + b3 + b4 == 4 and c1 + c2 == 2) else 1


if __name__ == "__main__":
    sys.exit(main())
