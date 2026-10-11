"""AMGT_CORPUS round 48, CRPS-ax-3 -- one read of whether the team's OTHER workbooks label any of the FactSet codes that still read "FactSet line"
unambiguously (CTO round-47 analysis s.3 item 3; "no inference"). Issuer's file. READ-ONLY: S3 GETs only.

  PYTHONPATH=. python ax3_label_read.py

TARGETS: the FactSet-prefix codes (FF_ FFI_ FG_ FE_ P_) that the STORED /13 `team_sheet` fields serve with `claim: absent` on at least one name.
CORPUS: every .xlsx / .xlsm under the `research-report` bronze prefix (all names, all files, all sheets) plus the stored master template.
For each sheet and each row whose FIRST cell is a target code (the data-feed rows keep the code in column A), the two rules of round 47, nothing else:
  metric_name   the row's column-B cell is a string that is not identifier-form: that string
  pass_through  the distinct labels (the first non-identifier string in the first four cells) of OTHER rows that carry, in some column, exactly the formula
                `=<same column letter><this row>` (optional `$`)
Prints per target code: the distinct labels with their kind and the number of workbooks that assign each; then the verdict per code:
  UNAMBIGUOUS   exactly one distinct label across every workbook that assigns one -> a candidate for a cited, measured label
  CONFLICT      two or more distinct labels
  NONE          no workbook assigns one
and the totals. A code that the team's own workbooks already measure per name (round 47) shows its labels here too: the read is of the CODE, across workbooks.
"""
import collections
import io
import re
import sys

import openpyxl

from lib.research.corpus import dossier, team_sheet
from lib.research.harvest import store

IDENT = re.compile(r"^[A-Z][A-Z0-9]*(_[A-Z0-9]+)+$")
FACTSET = ("FF_", "FFI_", "FG_", "FE_", "P_")
MAXR, MAXC = 400, 60


def _targets(client):
    out = collections.Counter()
    cur = dossier.METHOD_VERSION
    dossier.METHOD_VERSION = "amgt-corpus-dossier/13"
    try:
        for key in sorted(dossier._resolved_security_keys(client=client)):
            doc = dossier.read_dossier(key, client=client)["document"]
            ts = (doc.get("fields") or {}).get("team_sheet") or {}
            for lab, ent in zip(ts.get("rows") or [], ts.get("row_labels") or []):
                if ent and ent.get("source") == "factset" and ent.get("claim") == "absent" and ent.get("code", "").startswith(FACTSET):
                    out[ent["code"]] += 1
    finally:
        dossier.METHOD_VERSION = cur
    return out


def _read_sheet(ws_f, ws_v, targets, found, wb_name):
    """TWO views of one sheet, aligned by position: `ws_f` (formulas as text) decides WHICH rows point at a code row; `ws_v` (cached values) supplies the TEXT of
    every label, exactly as the round-47 reader does (a label cell that holds a formula is its VALUE, never the formula text: this file's second run listed
    '=+E101-1' as a label). Values only, row number = index + 1, column letter from the index: never cell attributes (an EmptyCell has none)."""
    from openpyxl.utils import get_column_letter

    rf = list(ws_f.iter_rows(min_row=1, max_row=MAXR, max_col=MAXC, values_only=True))
    rv = list(ws_v.iter_rows(min_row=1, max_row=MAXR, max_col=MAXC, values_only=True))
    n = min(len(rf), len(rv))
    for i in range(n):
        a = rv[i][0] if rv[i] else None
        if not (isinstance(a, str) and a.strip() in targets):
            continue
        code, rn = a.strip(), i + 1
        b = rv[i][1] if len(rv[i]) > 1 else None
        if isinstance(b, str) and b.strip() and not IDENT.match(b.strip()):
            found[code][(" ".join(b.split()), "metric_name")].add(wb_name)
        pat = re.compile(r"=\s*\$?([A-Z]{1,2})\$?%d$" % rn)
        for j in range(n):
            if j == i:
                continue
            lab = next((c for c in rv[j][:4] if isinstance(c, str) and c.strip() and not IDENT.match(c.strip())), None)
            if lab is None:
                continue
            for k, v in enumerate(rf[j]):
                if isinstance(v, str) and v.startswith("="):
                    m = pat.match(v.strip())
                    if m and m.group(1) == get_column_letter(k + 1):
                        found[code][(" ".join(lab.split()), "pass_through")].add(wb_name)
                        break


def _open(body):
    f = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    v = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    return f, v


def main():
    client = store._default_client()
    targets = _targets(client)
    print("targets: FactSet codes served with claim absent on at least one name: %d (rows: %d)" % (len(targets), sum(targets.values())))
    prefix = team_sheet._prefix()
    keys = [k for k in team_sheet._list_keys(client, prefix) if team_sheet._is_workbook(k)]
    tprefix = store.table_prefix(team_sheet.SOURCE, "bronze", "research_template")
    keys += [k for k in team_sheet._list_keys(client, tprefix) if team_sheet._is_workbook(k)]
    print("corpus: %d workbooks (%d research-report, %d template)" % (len(keys), len(keys) - len([k for k in keys if k.startswith(tprefix)]),
                                                                  len([k for k in keys if k.startswith(tprefix)])))
    found = collections.defaultdict(lambda: collections.defaultdict(set))
    opened = failed = 0
    sheets_ok = sheets_bad = 0
    for k in keys:
        try:
            body = client.get_object(Bucket=store.HARVEST_BUCKET, Key=k)["Body"].read()
            wb, wbv = _open(body)
        except Exception as exc:  # an unreadable workbook is counted and named, never dropped
            failed += 1
            print("UNREADABLE %s %s" % (k.rsplit("/", 1)[-1], str(exc)[:80]))
            continue
        opened += 1
        name = k.rsplit("/", 1)[-1]
        try:
            for ws in wb.worksheets:
                try:
                    _read_sheet(ws, wbv[ws.title], set(targets), found, name)
                    sheets_ok += 1
                except Exception as exc:  # counted and NAMED, never dropped: a swallowed error reads as an absence
                    sheets_bad += 1
                    print("SHEET-RAISED %s / %s: %s %s" % (name, ws.title, type(exc).__name__, str(exc)[:80]))
        finally:
            wb.close()
            wbv.close()
    print("workbooks opened %d failed %d of %d; sheets walked %d raised %d" % (opened, failed, len(keys), sheets_ok, sheets_bad))
    # POSITIVE CONTROL: the COPA workbook is known to label FF_SALES 'Revenue' (pass-through) and FE_ESTIMATE rows by their column-B metric name
    ctl = collections.defaultdict(lambda: collections.defaultdict(set))
    ck = [k for k in keys if "security_key=CPA_US_Equity" in k and k.lower().endswith(".xlsx") and "FactSet Based ONE PAGER" in k]
    if ck:
        cwb, cwv = _open(client.get_object(Bucket=store.HARVEST_BUCKET, Key=ck[0])["Body"].read())
        for ws in cwb.worksheets:
            _read_sheet(ws, cwv[ws.title], {"FF_SALES", "FE_ESTIMATE"}, ctl, "control")
        cwb.close()
        cwv.close()
    got_sales = sorted(lab for (lab, _k) in ctl.get("FF_SALES", {}))
    got_fe = sorted(lab for (lab, _k) in ctl.get("FE_ESTIMATE", {}))
    print("control COPA FF_SALES %s FE_ESTIMATE %s -> %s" % (got_sales, got_fe, "OK" if "Revenue" in got_sales and "Sales" in got_fe else "FAILED: the instrument does not read the known labels"))
    verdicts = collections.Counter()
    rows_by = collections.Counter()
    for code in sorted(targets):
        labs = found.get(code, {})
        distinct = sorted({lab for (lab, _kind) in labs})
        v = "NONE" if not distinct else ("UNAMBIGUOUS" if len(distinct) == 1 else "CONFLICT")
        verdicts[v] += 1
        rows_by[v] += targets[code]
        detail = "; ".join("%s [%s] x%d" % (lab, kind, len(wbs)) for (lab, kind), wbs in sorted(labs.items()))
        print("%-26s rows %4d  %-11s %s" % (code, targets[code], v, detail[:300]))
    print("read verdicts %s over %d codes; rows %s" % (dict(verdicts), len(targets), dict(rows_by)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
