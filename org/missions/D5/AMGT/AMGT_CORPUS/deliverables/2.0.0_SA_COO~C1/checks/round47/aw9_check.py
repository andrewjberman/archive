"""AMGT_CORPUS round 47, CRPS-aw-9 -- the master template is its own bronze table and the team's own labels are measured from the workbook.
Issuer's file. READ-ONLY (workbooks built in memory, a fake S3 client; no network).

  PYTHONPATH=. python aw9_check.py          prints  `aw9 build <b>/4 control <c>/2`

The rules under test (CTO round-46 ADDENDUM 2 msg 3994 and ADDENDUM 3 msg 4001):
  * `lib/research/corpus/template.py` owns the table `research_template` of the harvest bronze root. `put_template(body, item_id,
    client=None, source_mtime=None, file_name=None)` (file_name defaults to `team_sheet.MASTER_NAME`) stores the template with `store.bronze_put` (write-once, read-back verified) at partition
    `template=<item id>` and object `<first 16 hex of the sha256>/<file name>` through the address builder; it returns the bronze_put
    result dict (`key`, `put_performed`, `existed`, ...). A second put of the same bytes is a no-op (`existed` True).
  * the master arm of `team_sheet` (`rows_for`, `census`) reads that table: an object under `research_template/` whose basename is
    `team_sheet.MASTER_NAME` is the master; with none stored the answer is `master_absent` as today.
    A `master` answer carries NO team values: `row_values` and `row_formats` are all None (as long as `rows`) and `periods` is `[]`;
    the template's rows, row_classes, row_numbers, row_counts and row_labels are kept (AMENDMENT 3).
  * the team's own labels are MEASURED from the workbook being read, for a code row (identifier-form label) that no Bloomberg field-map
    description covers:
      - COLUMN-B METRIC NAME: when the code's row has a non-identifier string in column B (e.g. FE_ESTIMATE | "Sales"), the label is that
        text; `cite` names the sheet and the cell (`<sheet>!B<row>`);
      - SINGLE-ROW PASS-THROUGH: when exactly ONE distinct label (a non-identifier first string in the row's first four cells) belongs to
        rows whose period cell is exactly `=<column><code row>` (a direct reference to the code's own cell in the same column), the label
        is that text; `cite` names the sheet and the referencing cell (`<sheet>!D<row>`); two or more distinct labels leave the code as
        "FactSet line" (claim absent);
    the label dict is `{"text", "source": "factset", "claim": "measured", "cite", "code", "why": None}`. A code the field map describes
    keeps its Bloomberg description (claim held); a code with no measured label keeps round 47's fallback.
  * `python -m lib.research.corpus.template --labels <xlsx path>` prints the one line
    `template_labels measured <n> metric_name <m> pass_through <p> two_label_left <t> fallback <f> bloomberg_held <b> codes <c>` with
    n = m + p and c = n + t + f + b over the identifier-form code rows of the file's team sheet (t: a FactSet code that two or more
    distinct labels reference; f: a FactSet code with no measured label and not two-labelled; b: a code the field map describes).
"""
import contextlib
import datetime
import io
import sys
import tempfile

import openpyxl
from botocore.exceptions import ClientError

from lib.research.corpus import team_sheet
from lib.research.harvest import store
from lib.research.identity import crosswalk

YEARS = (2016, 2017, 2018, 2019, 2020)
COLS = "DEFGH"


class FakeS3:
    def __init__(self):
        self.objects = {}

    def put_object(self, Bucket=None, Key=None, Body=None, ContentType=None, Metadata=None, IfNoneMatch=None, **_kw):
        if IfNoneMatch == "*" and Key in self.objects:
            raise ClientError({"Error": {"Code": "PreconditionFailed", "Message": "exists"}}, "PutObject")
        self.objects[Key] = {"Body": bytes(Body), "Metadata": dict(Metadata or {}), "ContentType": ContentType,
                             "LastModified": datetime.datetime(2026, 10, 11, 16, 0, 0, tzinfo=datetime.timezone.utc)}
        return {}

    def head_object(self, Bucket=None, Key=None, **_kw):
        o = self.objects.get(Key)
        if o is None:
            raise ClientError({"Error": {"Code": "404", "Message": "nf"}}, "HeadObject")
        return {"Metadata": o["Metadata"], "ContentLength": len(o["Body"]), "LastModified": o["LastModified"]}

    def get_object(self, Bucket=None, Key=None, **_kw):
        return {"Body": io.BytesIO(self.objects[Key]["Body"])}

    def get_paginator(self, name):
        outer = self

        class _P:
            def paginate(self, Bucket=None, Prefix="", **_k):
                yield {"Contents": [{"Key": k} for k in sorted(outer.objects) if k.startswith(Prefix)]}

        return _P()


def workbook_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model NEW VERSION"
    ws["C3"] = "Financial Summary"
    for k, y in enumerate(YEARS):
        ws["%s3" % COLS[k]] = y
    # labelled rows that POINT at a code row (formulas: no cached value is written by openpyxl)
    for k, c in enumerate(COLS):
        ws["C4"] = "Revenue"
        ws["%s4" % c] = "=%s12" % c
        ws["C5"] = "Shares"
        ws["%s5" % c] = "=%s14" % c
        ws["C6"] = "Avg Diluted Shares"
        ws["%s6" % c] = "=%s14" % c
        ws["C7"] = "EBIT"
        ws["%s7" % c] = "=%s16" % c
    code_rows = [(12, "FF_SALES", None), (13, "FE_ESTIMATE", "Sales"), (14, "FF_COM_SHS_OUT_EPS_DIL", None),
                 (15, "FF_GROSS_INC", None), (16, "BEST_EBIT", None)]
    for r, code, param in code_rows:
        ws["A%d" % r] = code
        if param:
            ws["B%d" % r] = param
        for k, c in enumerate(COLS):
            ws["%s%d" % (c, r)] = 100 + 10 * r + k
    for i in range(6):
        ws["C%d" % (20 + i)] = "Filler %d" % i
        for k, c in enumerate(COLS):
            ws["%s%d" % (c, 20 + i)] = 5 + i + k
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    # Round 48 (CRPS-ax-5): the shipped factset_descriptions.json now holds nine cross-workbook entries, one of which (FF_COM_SHS_OUT_EPS_DIL) names a code this
    # fixture uses as its "two labels reference it" case. This check tests the MEASURED-label rules and the master arm, not the shipped data, so the entries
    # cache is emptied here (the entries' own precedence is ax5_check.py's).
    team_sheet._LABEL_CACHE["factset"] = {}
    body = workbook_bytes()
    client = FakeS3()
    b1 = b2 = b3 = b4 = 0
    # b1: the template table, write-once, through the builder
    try:
        from lib.research.corpus import template
        first = template.put_template(body, "01ITEMIDTEST", client=client)
        again = template.put_template(body, "01ITEMIDTEST", client=client)
        prefix = store.table_prefix("amgt", "bronze", "research_template")
        key = first.get("key", "")
        b1 = 1 if (key.startswith(prefix + "template=01ITEMIDTEST/") and key.endswith("/" + team_sheet.MASTER_NAME)
                   and first.get("put_performed") is True and again.get("existed") is True
                   and again.get("put_performed") is False) else 0
        print("aw9 b1 template put key %s; second put existed %s; ok %s" % (key, again.get("existed"), b1 == 1))
    except Exception as exc:  # noqa: BLE001
        print("aw9 b1 threw: %s" % str(exc)[:200])
        template = None
    # the template put above used the file name of the master only if put_template names it so: use the master's name explicitly
    # b2: the master arm reads that table (no research-report objects exist at all)
    try:
        client2 = FakeS3()
        if template is not None:
            template.put_template(body, "01ITEMIDTEST", client=client2, file_name=team_sheet.MASTER_NAME)
        got = team_sheet.rows_for("NOBODY_US_Equity", client=client2)
        # AMENDMENT 3 (issuer, 2026-10-11 ~19:4xZ): the master template's cells are ANOTHER company's figures (the template is filled with
        # one sample company's numbers); a master answer therefore carries NO values, formats or periods, only the template's rows, classes
        # and labels. Read on the live store: 63 names x 126 FIGURE rows would have shown that company's numbers as "team model".
        rv = got.get("row_values") or []
        rf = got.get("row_formats") or []
        nvals = sum(1 for v in rv if v is not None)
        nfmts = sum(1 for v in rf if v is not None)
        b2 = 1 if (got.get("class") == "master" and "research_template" in str(got.get("workbook_key"))
                   and "FF_SALES" in (got.get("rows") or [])
                   and len(rv) == len(got.get("rows") or []) and nvals == 0
                   and len(rf) == len(got.get("rows") or []) and nfmts == 0
                   and got.get("periods") == []
                   and "FIGURE" in (got.get("row_classes") or [])) else 0
        print("aw9 b2 master arm class %s workbook_key %s; values %d formats %d periods %s; ok %s"
              % (got.get("class"), got.get("workbook_key"), nvals, nfmts, got.get("periods"), b2 == 1))
    except Exception as exc:  # noqa: BLE001
        print("aw9 b2 threw: %s" % str(exc)[:200])
    # b3: measured labels from the own workbook being read
    try:
        client3 = FakeS3()
        own_key = (team_sheet._prefix() + crosswalk.security_key_partition("X_US_Equity")
                   + "/abc123/X FactSet Based ONE PAGER 2026 10 11_v1.xlsx")
        client3.put_object(Key=own_key, Body=body, Metadata={"source-mtime": "2026-10-11T00:00:00Z"})
        got = team_sheet.rows_for("X_US_Equity", client=client3)
        by = {}
        for lab, ent in zip(got.get("rows") or [], got.get("row_labels") or []):
            by.setdefault(lab, ent)
        s = by.get("FF_SALES") or {}
        fe = by.get("FE_ESTIMATE") or {}
        two = by.get("FF_COM_SHS_OUT_EPS_DIL") or {}
        gi = by.get("FF_GROSS_INC") or {}
        be = by.get("BEST_EBIT") or {}
        b3 = 1 if (s.get("text") == "Revenue" and s.get("claim") == "measured" and s.get("source") == "factset"
                   and "Model NEW VERSION" in str(s.get("cite")) and "D4" in str(s.get("cite"))
                   and fe.get("text") == "Sales" and fe.get("claim") == "measured" and "B13" in str(fe.get("cite"))
                   and two.get("text") == "FactSet line" and two.get("claim") == "absent"
                   and gi.get("text") == "FactSet line" and gi.get("claim") == "absent"
                   and be.get("claim") == "held" and be.get("source") == "bloomberg") else 0
        print("aw9 b3 labels FF_SALES %r FE_ESTIMATE %r two-label %r no-ref %r Bloomberg %r; ok %s"
              % (s.get("text"), fe.get("text"), two.get("text"), gi.get("text"), be.get("claim"), b3 == 1))
    except Exception as exc:  # noqa: BLE001
        print("aw9 b3 threw: %s" % str(exc)[:200])
    # b4: the CLI count line
    try:
        from lib.research.corpus import template as tmod
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
            fh.write(body)
            path = fh.name
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            tmod.main(["--labels", path])
        line = [x for x in out.getvalue().splitlines() if x.startswith("template_labels ")]
        want = "template_labels measured 2 metric_name 1 pass_through 1 two_label_left 1 fallback 1 bloomberg_held 1 codes 5"
        b4 = 1 if line and line[-1].strip() == want else 0
        print("aw9 b4 line %r; ok %s" % (line[-1] if line else None, b4 == 1))
    except Exception as exc:  # noqa: BLE001
        print("aw9 b4 threw: %s" % str(exc)[:200])
    # control c1: with no template and no own workbook the answer is master_absent with every class key empty
    c1 = 0
    try:
        got = team_sheet.rows_for("NOBODY_US_Equity", client=FakeS3())
        c1 = 1 if (got.get("class") == "master_absent" and got.get("rows") == [] and got.get("row_classes") == []) else 0
        print("aw9 c1 master_absent with nothing stored %s" % (c1 == 1))
    except Exception as exc:  # noqa: BLE001
        print("aw9 c1 threw: %s" % str(exc)[:200])
    # control c2: the round-46 class reading of the workbook is unchanged (rows, classes, values of the period cells)
    c2 = 0
    try:
        rd = team_sheet.sheet_reading(body).get("Model NEW VERSION") or {}
        rows = rd.get("rows") or []
        vals = rd.get("row_values") or []
        i = rows.index("FF_SALES")
        c2 = 1 if (rd.get("periods") and [p.get("label") for p in rd["periods"]] == [str(y) for y in YEARS]
                   and vals[i] == [220, 221, 222, 223, 224]
                   and rd["row_classes"][i] == "FIGURE") else 0
        print("aw9 c2 round-46 reading unchanged %s" % (c2 == 1))
    except Exception as exc:  # noqa: BLE001
        print("aw9 c2 threw: %s" % str(exc)[:200])
    print("aw9 build %d/4 control %d/2" % (b1 + b2 + b3 + b4, c1 + c2))
    return 0 if (b1 + b2 + b3 + b4 == 4 and c1 + c2 == 2) else 1


if __name__ == "__main__":
    sys.exit(main())
