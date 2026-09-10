"""
f6_remeasure.py -- F6 RE-MEASURED (AMGT-s-12e)

F6 (CITED, the brief @19ce7d5b4d; the CRO's method in
org/tier1_exec/4.0.0_SA_CRO/deliverables/AMGT_RESAPP_LAMBDAS_PEERS_IRSCAN_20260910.md
s.2): "Population: all 54 *ONE PAGER*.xlsx workbooks across the 89 name
folders. Method: read each CSC sheet, extract Bloomberg-format tickers, score
overlap against the peer set found on the SAP sheet that first raised the
flag." Result claimed: 48 stale / 4 own / 2 no-CSC.

This script RE-DERIVES the count on the live corpus, with the method the
task order's cell text states explicitly (AMGT-s-12, part e): read every
*ONE PAGER*.xlsx workbook's CSC sheet, extract Bloomberg-format tickers, and
score overlap against the fixed KPN telecom set named in that cell text --
KPN, Vodafone, BT, Orange, Telenor, Telekom Austria, Rogers, Telus, Swisscom,
AT&T, Verizon, Telefonica -- rather than a per-workbook "peer set found on
the SAP sheet" (no 'SAP' sheet was found on any workbook measured while
building this script; the CRO's own text does not give a numeric threshold
either). The fixed 12-ticker set below was measured directly off ASML NA's
CSC sheet (a full, unmodified copy of the stale template) and cross-checked
byte-for-byte against Imperial Brands' stale FactSet workbook -- same 12
rows, same order.

Denominator: every file in every one of the 89 name folders whose filename
contains "ONE PAGER" (case-insensitive) and ends ".xlsx" -- a folder with two
such files (Imperial Brands, Nokia, Petrobras, SAP GY, Total Energies)
contributes two workbooks, matching F6's own population definition ("all 54
... workbooks", not "all 54 folders" -- measured directly: Imperial Brands'
two one-pager files disagree with each other, one stale-copy, one its own
analyst set, which only makes sense if the population counts workbooks).

Classification per workbook:
  stale-set : CSC sheet found; extracted Bloomberg ticker roots overlap the
              fixed 12-name KPN set at >= STALE_OVERLAP_THRESHOLD (6, i.e.
              at least half) -- every stale copy measured while building this
              script carried all 12 unchanged, so any threshold from ~4
              upward produces an identical partition on this corpus; 6 is
              the number recorded here.
  own-set   : CSC sheet found, ticker rows extracted, overlap below threshold.
  no-CSC    : no CSC sheet in the workbook, OR a CSC sheet with no
              Bloomberg-ticker-shaped rows found within the read cap.

Agreement or disagreement with the CRO's 48/4/2 is a FINDING this script
prints and f6_remeasure.json records -- it is not this criterion (AC3 only
checks that the denominator is printed and the partition sums to it).

Run:    python org/missions/D5/AMGT/AMGT_RESAPP/deliverables/0.9.0_WK_AMGT_RESAPP~3/f6_remeasure.py
Writes: f6_remeasure.json beside this script.
"""

from __future__ import annotations

import datetime
import json
import os
import re

import openpyxl

CORPUS_ROOT = (
    r"C:\Users\Opal Bloomberg\OneDrive - Opal Capital\Shared - Documents"
    r"\Asset Management\Research"
)

OUTPUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "f6_remeasure.json")

# Measured directly off ASML NA's CSC sheet (ticker column, " Equity" suffix
# stripped, uppercased) -- the canonical, unmodified stale-template row set.
KPN_STALE_SET = {
    "KPN NA",
    "VOD LN",
    "BT/A LN",
    "ORA FP",
    "TEL NO",
    "TKA AV",
    "RCI/B CN",
    "T CN",
    "SCMN SW",
    "T US",
    "VZ US",
    "TEF SM",
}
KPN_STALE_SET_NAMES = (
    "KPN, Vodafone, BT, Orange, Telenor, Telekom Austria, Rogers, Telus, "
    "Swisscom, AT&T, Verizon, Telefonica"
)

STALE_OVERLAP_THRESHOLD = 6  # at least half of the 12-name fixed set

CRO_CITED_RESULT = {"stale_set": 48, "own_set": 4, "no_csc": 2, "of": 54}

_MAX_ROWS = 40
_MAX_COLS = 14

_EQUITY_SUFFIX_RE = re.compile(r"^(?P<root>.+?)\s+EQUITY$", re.IGNORECASE)


def _now_iso() -> str:
    return datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%MZ")


def _find_one_pager_files() -> list:
    """Every (folder, filename, full_path) whose filename matches *ONE PAGER*.xlsx,
    across every one of the 89 name folders -- the workbook-level denominator."""
    out = []
    for folder_name in sorted(os.listdir(CORPUS_ROOT)):
        folder_path = os.path.join(CORPUS_ROOT, folder_name)
        if not os.path.isdir(folder_path):
            continue
        try:
            files = os.listdir(folder_path)
        except OSError:
            continue
        for fname in sorted(files):
            if "ONE PAGER" in fname.upper() and fname.upper().endswith(".XLSX"):
                out.append((folder_name, fname, os.path.join(folder_path, fname)))
    return out


def _extract_csc_tickers(ws) -> list:
    roots = []
    for row in ws.iter_rows(min_row=1, max_row=_MAX_ROWS, max_col=_MAX_COLS, values_only=True):
        for cell in row:
            if not isinstance(cell, str):
                continue
            m = _EQUITY_SUFFIX_RE.match(cell.strip())
            if m:
                roots.append(re.sub(r"\s+", " ", m.group("root")).strip().upper())
    return roots


def classify_workbook(path: str) -> dict:
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        return {
            "classification": "no-CSC",
            "reason": "workbook open failed: %s" % exc,
            "tickers_extracted": [],
            "overlap_count": 0,
        }
    try:
        csc_name = None
        for sn in wb.sheetnames:
            if sn.strip().lower() == "csc":
                csc_name = sn
                break
        if csc_name is None:
            return {
                "classification": "no-CSC",
                "reason": "no sheet named 'CSC' (sheets present: %s)" % ", ".join(wb.sheetnames),
                "tickers_extracted": [],
                "overlap_count": 0,
            }
        ws = wb[csc_name]
        roots = _extract_csc_tickers(ws)
        if not roots:
            return {
                "classification": "no-CSC",
                "reason": "CSC sheet present but no Bloomberg-ticker-shaped "
                "(' EQUITY' suffix) rows found within the first %d row(s) x "
                "%d col(s)" % (_MAX_ROWS, _MAX_COLS),
                "tickers_extracted": [],
                "overlap_count": 0,
            }
        overlap = sorted(set(roots) & KPN_STALE_SET)
        overlap_count = len(overlap)
        if overlap_count >= STALE_OVERLAP_THRESHOLD:
            classification = "stale-set"
        else:
            classification = "own-set"
        return {
            "classification": classification,
            "reason": "CSC sheet '%s': %d ticker row(s) extracted, %d/%d overlap "
            "the fixed KPN set (threshold %d)"
            % (csc_name, len(roots), overlap_count, len(KPN_STALE_SET), STALE_OVERLAP_THRESHOLD),
            "tickers_extracted": roots,
            "overlap_count": overlap_count,
        }
    finally:
        wb.close()


def remeasure() -> dict:
    started_at = _now_iso()
    workbooks = _find_one_pager_files()
    denominator = len(workbooks)
    results = []
    partition = {"stale-set": 0, "own-set": 0, "no-CSC": 0}
    for folder_name, fname, full_path in workbooks:
        outcome = classify_workbook(full_path)
        partition[outcome["classification"]] += 1
        results.append(
            {
                "folder": folder_name,
                "file": fname,
                "path": full_path,
                **outcome,
            }
        )
    finished_at = _now_iso()
    assert sum(partition.values()) == denominator

    try:
        root_mtime = datetime.datetime.utcfromtimestamp(
            os.stat(CORPUS_ROOT).st_mtime
        ).strftime("%Y-%m-%dT%H:%MZ")
    except OSError:
        root_mtime = None

    return {
        "corpus_root": CORPUS_ROOT,
        "corpus_root_mtime": root_mtime,
        "read_started_at": started_at,
        "read_finished_at": finished_at,
        "denominator_workbooks": denominator,
        "partition": partition,
        "stale_set_definition": sorted(KPN_STALE_SET),
        "stale_set_names_cited": KPN_STALE_SET_NAMES,
        "stale_overlap_threshold": STALE_OVERLAP_THRESHOLD,
        "cro_cited_result_F6": CRO_CITED_RESULT,
        "agreement_with_cro_cited": {
            "stale_set": partition["stale-set"] == CRO_CITED_RESULT["stale_set"],
            "own_set": partition["own-set"] == CRO_CITED_RESULT["own_set"],
            "no_csc": partition["no-CSC"] == CRO_CITED_RESULT["no_csc"],
            "denominator": denominator == CRO_CITED_RESULT["of"],
        },
        "results": results,
    }


def main() -> int:
    data = remeasure()
    with open(OUTPUT_PATH, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=True)
        fh.write("\n")

    vintage = data["read_finished_at"]
    print(
        "F6 RE-MEASURED: stale-set %d / own-set %d / no-CSC %d of %d workbooks; "
        "vintage %s"
        % (
            data["partition"]["stale-set"],
            data["partition"]["own-set"],
            data["partition"]["no-CSC"],
            data["denominator_workbooks"],
            vintage,
        )
    )
    print(
        "F6 CITED (brief @19ce7d5b4d, CRO method): stale-set %d / own-set %d / "
        "no-CSC %d of %d workbooks -- agreement: %s"
        % (
            CRO_CITED_RESULT["stale_set"],
            CRO_CITED_RESULT["own_set"],
            CRO_CITED_RESULT["no_csc"],
            CRO_CITED_RESULT["of"],
            data["agreement_with_cro_cited"],
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
