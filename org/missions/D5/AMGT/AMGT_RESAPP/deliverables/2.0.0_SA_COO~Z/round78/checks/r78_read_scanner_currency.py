"""AMGT_RESAPP round 78 -- READ: does the scanner's own source carry a currency (CTO round-78 ruling on bb-8).

  python r78_read_scanner_currency.py [--write <folder>] [--skip-served]

A READ. No build, no store write, no vendor call. It reads:
  - the six Security Scanner universe files the peer rows are built from
    (s3://opal-data-portal-v2/json/security_scanner/latest/universes/<id>/universe_data.json; reader
    lib/research/peers/universe.py);
  - the UPSTREAM source those files are built from: the FactSet screening workbook the producer
    (infra/aws/lambda/opal_security_scanner_processor/lambda_function.py, lambda_handler) reads with
    pd.read_excel(header=4), under s3://opal-data-portal-v2/security_scanner/uploads/<asOf>/. The workbook whose
    folder equals the served files' asOf is the one read; its four report-header rows are read too;
  - unless --skip-served: one GET-shaped direct invoke of the deployed function's /overlay route per covered name, to
    count the peer rows that carry a price and no price currency today.

A column NAMES A CURRENCY when its header matches CURRENCY_NAME (currency, ccy, crncy, fx, iso, denom). Headers that
only begin "curr" (curr_ratio, curr_asset_liq: the CURRENT ratio and CURRENT-asset liquidity factors) are printed apart
with their values, which are numbers; they are not currencies. A column also counts when its VALUES are currency codes
(three capital letters, the set including USD), whatever its header. The verdict on the question is the count of such
columns in the workbook, and the count of served peer rows a pass-through would reach (0 when there is none).

One further join is measured, because the workbook carries it and the producer drops it: each row's SEDOL and ISIN.
For the served peer rows with no price currency, how many have a workbook SEDOL or ISIN equal to a pool-list record's
stored SEDOL or ISIN. That is a count of rows a listing-exact match could reach, not a currency.
Verdict: R78 READ SCANNER CURRENCY: DONE (exit 0) when the six files and the workbook were both read.
"""
from __future__ import annotations

import argparse
import collections
import io
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor

SCANNER_BUCKET = "opal-data-portal-v2"
SCANNER_PREFIX = "json/security_scanner/latest/universes/"
UPLOAD_PREFIX = "security_scanner/uploads/"
FN = "opal-amgt-scoped-api"
PREFIX = "/api/scoped/amgt"
POOL_BUCKET = "opal-data-harvest"
# The six universes the producer writes (its UNIVERSE_MAP) and the peer engine reads.
UNIVERSES = ("sp500", "r1000", "r1000v", "acwx", "eafe", "eafev")
CURRENCY_NAME = re.compile(r"currenc|ccy|crncy|(?<![a-z])fx(?![a-z])|(?<![a-z])iso(?![a-z])|denom", re.IGNORECASE)
CURR_PREFIX = re.compile(r"^curr", re.IGNORECASE)
LISTING_COLUMNS = ("price", "price_split", "price_unadj", "country_name", "country_exch", "country_hq",
                   "country_incorp", "country_risk", "adr_flag", "primary_list", "is_primary", "isin", "sedol", "cusip",
                   "mkt_val")


def _event(path):
    return {"rawPath": path, "requestContext": {"http": {"method": "GET"}}, "headers": {"x-firm-id": "opal"},
            "queryStringParameters": None}


def _list(s3, bucket, prefix):
    out = []
    for page in s3.get_paginator("list_objects_v2").paginate(Bucket=bucket, Prefix=prefix):
        out += page.get("Contents", []) or []
    return out


def read_scanner_files(s3):
    files = {}
    for obj in _list(s3, SCANNER_BUCKET, SCANNER_PREFIX):
        if not obj["Key"].endswith("/universe_data.json"):
            continue
        universe_id = obj["Key"][len(SCANNER_PREFIX):].split("/", 1)[0]
        if universe_id not in UNIVERSES:
            files.setdefault("_other_folders", []).append(
                {"id": universe_id, "stored": obj["LastModified"].strftime("%Y-%m-%dT%H:%MZ")})
            continue
        doc = json.loads(s3.get_object(Bucket=SCANNER_BUCKET, Key=obj["Key"])["Body"].read())
        rows = doc.get("securities") or []
        fields = collections.Counter(k for r in rows for k in r)
        files[universe_id] = {
            "key": obj["Key"], "stored": obj["LastModified"].strftime("%Y-%m-%dT%H:%MZ"), "as_of": doc.get("asOf"),
            "top_level_keys": sorted(doc.keys()), "rows": len(rows), "row_fields": dict(fields),
            "rows_with_price": sum(1 for r in rows if r.get("price") is not None),
            "value_codes": len({c for r in rows for c in (r.get("values") or {})}),
            "currency_named_row_fields": sorted(k for k in fields if CURRENCY_NAME.search(k)),
            "currency_named_value_codes": sorted({c for r in rows for c in (r.get("values") or {})
                                                  if CURRENCY_NAME.search(c)}),
            "currency_named_top_level": sorted(k for k in doc if CURRENCY_NAME.search(k)),
        }
    others = files.pop("_other_folders", [])
    return files, others


def read_workbook(s3, as_of):
    import openpyxl
    import pandas as pd

    uploads = [o for o in _list(s3, SCANNER_BUCKET, UPLOAD_PREFIX) if o["Key"].lower().endswith(".xlsx")]
    folders = sorted({o["Key"][len(UPLOAD_PREFIX):].split("/", 1)[0] for o in uploads})
    match = [o for o in uploads if o["Key"][len(UPLOAD_PREFIX):].split("/", 1)[0] == as_of]
    chosen = sorted(match or uploads, key=lambda o: o["LastModified"])[-1]
    body = s3.get_object(Bucket=SCANNER_BUCKET, Key=chosen["Key"])["Body"].read()
    book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    sheets = book.sheetnames
    sheet = book[sheets[0]]
    head_rows = []
    for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
        head_rows.append([str(c) for c in row if c is not None][:400])
    book.close()
    frame = pd.read_excel(io.BytesIO(body), header=4)
    columns = [str(c) for c in frame.columns]
    named = [c for c in columns if CURRENCY_NAME.search(c)]
    curr_prefixed = [c for c in columns if CURR_PREFIX.search(c) and c not in named]
    identifiers = {}
    for _, row in frame.iterrows():
        ticker = str(row.get("ticker", ""))
        if ticker and ticker not in identifiers:
            identifiers[ticker] = {"sedol": None if pd.isna(row.get("sedol")) else str(row.get("sedol")),
                                   "isin": None if pd.isna(row.get("isin")) else str(row.get("isin")),
                                   "country_exch": None if pd.isna(row.get("country_exch")) else str(
                                       row.get("country_exch"))}
    listing = {}
    for column in LISTING_COLUMNS:
        if column in frame.columns:
            series = frame[column]
            listing[column] = {"non_null": int(series.notna().sum()), "distinct": int(series.nunique(dropna=True)),
                               "sample": [str(v) for v in series.dropna().unique()[:8]]}
    report_text = " | ".join(" ".join(r) for r in head_rows[:4])
    return {
        "key": chosen["Key"], "stored": chosen["LastModified"].strftime("%Y-%m-%dT%H:%MZ"), "bytes": chosen["Size"],
        "folder_matches_served_as_of": bool(match), "upload_folders": folders, "sheets": sheets,
        "report_header_rows": [r[:12] for r in head_rows[:4]], "report_header_mentions_currency":
            bool(re.search(r"currenc|\bUSD\b|\blocal\b", report_text, re.IGNORECASE)),
        "rows": int(len(frame)), "columns": len(columns), "column_names": columns,
        "currency_named_columns": named,
        "curr_prefixed_columns": {c: {"non_null": int(frame[c].notna().sum()), "numeric": bool(
            pd.api.types.is_numeric_dtype(frame[c])), "sample": [str(v) for v in frame[c].dropna().unique()[:4]]}
            for c in curr_prefixed},
        "identifiers_by_ticker": identifiers,
        "currency_named_columns_detail": {c: {"non_null": int(frame[c].notna().sum()),
                                              "sample": [str(v) for v in frame[c].dropna().unique()[:10]]}
                                          for c in named},
        "listing_columns": listing,
        "text_columns_with_currency_code_values": _code_valued_columns(frame),
    }


def _code_valued_columns(frame):
    """Columns whose values look like currency codes (three capital letters, a small distinct set that includes USD)."""
    found = {}
    for column in frame.columns:
        series = frame[column].dropna()
        if series.empty or series.dtype != object:
            continue
        values = set(str(v) for v in series.unique()[:400])
        if "USD" in values and all(re.fullmatch(r"[A-Z]{3}", v) for v in values):
            found[str(column)] = sorted(values)[:20]
    return found


def read_served_peers():
    import boto3

    from lib.research.coverage import covered_universe

    lam = boto3.client("lambda", region_name="us-east-1")
    keys = sorted(covered_universe.covered_security_keys())

    def one(key):
        raw = json.loads(lam.invoke(FunctionName=FN, Payload=json.dumps(
            _event("%s/security/%s/overlay" % (PREFIX, key))).encode())["Payload"].read())
        try:
            body = json.loads(raw.get("body") or "{}")
        except (ValueError, TypeError):
            body = {}
        members = (body.get("peers") or {}).get("members") or []
        return raw.get("statusCode"), [m for m in members if isinstance(m, dict)]

    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(one, keys))
    rows = [m for _status, members in results for m in members]
    priced = [m for m in rows if m.get("price") is not None]
    return {
        "covered": len(keys), "not_200": sum(1 for status, _m in results if status != 200),
        "names_with_a_peer_set": sum(1 for _status, members in results if members),
        "peer_rows": len(rows), "peer_rows_with_price": len(priced),
        "priced_with_price_currency": sum(1 for m in priced if m.get("price_currency")),
        "priced_without_price_currency": sum(1 for m in priced if not m.get("price_currency")),
        "no_currency_rows": [{"ticker": str(m.get("ticker")), "country": m.get("country")} for m in priced
                             if not m.get("price_currency")],
        "member_fields": sorted({k for m in rows for k in m}),
        "currency_named_member_fields": sorted({k for m in rows for k in m if CURRENCY_NAME.search(k)}),
    }


def build_table(skip_served):
    import boto3

    s3 = boto3.client("s3", region_name="us-east-1")
    files, others = read_scanner_files(s3)
    as_ofs = sorted({f["as_of"] for f in files.values() if f["as_of"]})
    workbook = read_workbook(s3, as_ofs[-1] if as_ofs else "")
    served = None if skip_served else read_served_peers()
    union_fields = sorted({k for f in files.values() for k in f["row_fields"]})
    join = None
    if served is not None:
        pool_keys = sorted(o["Key"] for o in _list(s3, POOL_BUCKET, "crosswalk/v1/") if o["Key"].endswith(".json"))
        pool = json.loads(s3.get_object(Bucket=POOL_BUCKET, Key=pool_keys[-1])["Body"].read())
        sedols = {str((r.get("identity") or {}).get("sedol")): r["security_key"] for r in pool["records"]
                  if (r.get("identity") or {}).get("sedol")}
        isins = {str((r.get("identity") or {}).get("isin")): r["security_key"] for r in pool["records"]
                 if (r.get("identity") or {}).get("isin")}
        ids = workbook["identifiers_by_ticker"]
        rows = served.pop("no_currency_rows")
        in_book = [r for r in rows if r["ticker"] in ids]
        matched = [r for r in in_book if ids[r["ticker"]]["sedol"] in sedols or ids[r["ticker"]]["isin"] in isins]
        join = {"pool_list": pool_keys[-1], "pool_records": len(pool["records"]),
                "pool_records_with_a_sedol": len(sedols), "pool_records_with_an_isin": len(isins),
                "rows_without_currency": len(rows), "of_those_in_the_workbook_by_ticker": len(in_book),
                "of_those_with_a_sedol_or_isin_in_the_workbook": sum(
                    1 for r in in_book if ids[r["ticker"]]["sedol"] or ids[r["ticker"]]["isin"]),
                "of_those_equal_to_a_pool_record": len(matched),
                "matched_distinct_tickers": sorted({r["ticker"] for r in matched}),
                "rows_without_currency_by_country": dict(collections.Counter(
                    "United States" if r["country"] == "United States" else "other" for r in rows))}
    workbook_ids = workbook.pop("identifiers_by_ticker")
    workbook["distinct_tickers"] = len(workbook_ids)
    carries = bool(workbook["currency_named_columns"] or workbook["text_columns_with_currency_code_values"])
    gain = None
    if served is not None:
        gain = served["priced_without_price_currency"] if carries else 0
    return {
        "scanner_files": files, "scanner_rows": sum(f["rows"] for f in files.values()),
        "scanner_row_fields": union_fields, "scanner_as_of": as_ofs,
        "scanner_currency_named_fields": sorted({k for f in files.values() for k in
                                                 f["currency_named_row_fields"] + f["currency_named_value_codes"]
                                                 + f["currency_named_top_level"]}),
        "workbook": workbook, "upstream_carries_a_currency": carries, "served_peers": served,
        "other_universe_folders": others, "identifier_join": join,
        "rows_that_would_gain_a_currency": gain,
    }


def print_table(t):
    print("THE SCANNER FILES (%d files, %d rows, asOf %s)" % (len(t["scanner_files"]), t["scanner_rows"],
                                                              t["scanner_as_of"]))
    for universe_id, f in sorted(t["scanner_files"].items()):
        print("   %-7s rows %5d with a price %5d; factor codes %d; stored %s; row fields %s"
              % (universe_id, f["rows"], f["rows_with_price"], f["value_codes"], f["stored"], sorted(f["row_fields"])))
    print("   row fields over the six files: %d %s" % (len(t["scanner_row_fields"]), t["scanner_row_fields"]))
    print("   fields, factor codes or top-level keys naming a currency: %d %s"
          % (len(t["scanner_currency_named_fields"]), t["scanner_currency_named_fields"]))
    print("   other folders under the same prefix, not written by the producer today and not read here: %s"
          % t["other_universe_folders"])
    w = t["workbook"]
    print("THE UPSTREAM SOURCE (the producer's input)")
    print("   s3://%s/%s; %d bytes; stored %s; its folder equals the served asOf: %s; upload folders %s"
          % (SCANNER_BUCKET, w["key"], w["bytes"], w["stored"], w["folder_matches_served_as_of"], w["upload_folders"]))
    print("   sheets %s; data rows %d; columns %d" % (w["sheets"], w["rows"], w["columns"]))
    for index, row in enumerate(w["report_header_rows"], start=1):
        print("   report header row %d: %s" % (index, row))
    print("   the report header mentions a currency: %s" % w["report_header_mentions_currency"])
    print("   columns whose NAME names a currency: %d of %d %s" % (len(w["currency_named_columns"]), w["columns"],
                                                                  json.dumps(w["currency_named_columns_detail"])))
    print("   columns whose name only begins 'curr' (current-ratio factors, numbers, not currencies): %s"
          % json.dumps(w["curr_prefixed_columns"]))
    print("   text columns whose VALUES are currency codes: %d %s"
          % (len(w["text_columns_with_currency_code_values"]), json.dumps(w["text_columns_with_currency_code_values"])))
    print("   the listing columns the workbook does carry:")
    for column, detail in w["listing_columns"].items():
        print("      %-15s non-null %5d of %d; distinct %5d; e.g. %s" % (column, detail["non_null"], w["rows"],
                                                                        detail["distinct"], detail["sample"][:6]))
    print("   all %d column names: %s" % (w["columns"], ", ".join(w["column_names"])))
    s = t["served_peers"]
    if s is not None:
        print("THE SERVED PEER ROWS (deployed /overlay, %d covered names, not 200: %d)" % (s["covered"], s["not_200"]))
        print("   names with a peer set %d; peer rows %d; with a price %d; of those with a price currency %d, without %d"
              % (s["names_with_a_peer_set"], s["peer_rows"], s["peer_rows_with_price"],
                 s["priced_with_price_currency"], s["priced_without_price_currency"]))
    j = t["identifier_join"]
    if j is not None:
        print("THE IDENTIFIER JOIN (the workbook's SEDOL and ISIN, which the producer does not pass through)")
        print("   rows with no price currency %d (%s); in the workbook by ticker %d; with a SEDOL or ISIN there %d; "
              "equal to a pool-list record's stored SEDOL or ISIN %d %s"
              % (j["rows_without_currency"], j["rows_without_currency_by_country"],
                 j["of_those_in_the_workbook_by_ticker"], j["of_those_with_a_sedol_or_isin_in_the_workbook"],
                 j["of_those_equal_to_a_pool_record"], j["matched_distinct_tickers"]))
        print("   pool list %s: %d records, %d with a stored SEDOL, %d with a stored ISIN"
              % (j["pool_list"], j["pool_records"], j["pool_records_with_a_sedol"], j["pool_records_with_an_isin"]))
    print("THE ANSWER: the upstream source carries a currency per row: %s; peer rows that would gain a currency from "
          "a pass-through: %s" % ("YES" if t["upstream_carries_a_currency"] else "NO",
                                  t["rows_that_would_gain_a_currency"]))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--write", default=None, help="folder to write SCANNER_CURRENCY_READ.json to")
    ap.add_argument("--skip-served", action="store_true", help="do not read the deployed /overlay route")
    a = ap.parse_args()
    table = build_table(a.skip_served)
    print_table(table)
    if a.write:
        os.makedirs(a.write, exist_ok=True)
        with open(os.path.join(a.write, "SCANNER_CURRENCY_READ.json"), "w", encoding="utf-8", newline="\n") as fh:
            json.dump(table, fh, indent=1, sort_keys=True)
            fh.write("\n")
        print("   wrote SCANNER_CURRENCY_READ.json to %s" % a.write)
    ok = len(table["scanner_files"]) == 6 and table["workbook"]["columns"] > 0
    print("R78 READ SCANNER CURRENCY: %s" % ("DONE" if ok else "NOT DONE"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
