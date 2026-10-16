"""ADVI_COTBKF TO-CB-da-1 (issuer's check): is the Cottonwood account-match review workbook readable by a person and complete?

Opens the workbook the worker wrote (default: the desktop file), the committed packet (to learn the 107 pairs to confirm) and the OFF-GIT key file (to
learn every account string, so the check can prove none is visible). Prints counts only; no account, name or value.
  python cz1_review_sheet_check.py [workbook.xlsx]
    REVIEW-SHEET sheets_ok=<bool> rows=<n> ids_match=<bool> link_id_hidden_last=<bool> answer_blank=<bool> answer_validation=<bool>
                 visible_account_strings=<n> visible_link_ids=<n> blank_old_name=<n> blank_new_name=<n> blank_value=<n> strengths=<s:n,...>
  python cz1_review_sheet_check.py --offgit [workbook.xlsx] <record.json> <record.md>
    SHEET-OFFGIT outside_repo=<bool> tracked_review_xlsx=<n> record_account_strings=<n>
Rules: sheets_ok = the first sheet is named 'Read me' with at least 600 characters of text and a sheet named 'To confirm' exists; rows = data rows of
'To confirm'; ids_match = the hidden last column holds exactly the packet's to-confirm link ids, once each; answer_blank = the column headed 'Your answer'
is empty in every row; answer_validation = that column carries a Yes/No/Unsure list validation; visible_account_strings = visible cells (not the hidden last
column) that equal a full account string of the key; visible_link_ids = visible cells that equal a link id; strengths counts the 'Match strength' column.
Needs openpyxl and the off-git key file at C:/dev/_cotbkf_match_20261015/NEAR_MATCH_REVIEW_KEY.csv.
"""
import csv
import subprocess
import sys
from collections import Counter
from pathlib import Path

import openpyxl

KEY = "C:/dev/_cotbkf_match_20261015/NEAR_MATCH_REVIEW_KEY.csv"
PACKET = "org/missions/D8/ADVI/ADVI_COTBKF/deliverables/0.9.0_WK_ADVI_COTBKF~180/NEAR_MATCH_REVIEW_PACKET.csv"
DEFAULT = "C:/Users/Test/Desktop/Cottonwood_Account_Match_Review_20261016.xlsx"
REPO = Path(subprocess.run(["git", "rev-parse", "--show-toplevel"], capture_output=True, text=True).stdout.strip() or ".").resolve()

args = [a for a in sys.argv[1:] if not a.startswith("--")]
offgit = "--offgit" in sys.argv
wb_path = args[0] if args and args[0].lower().endswith(".xlsx") else DEFAULT

key = list(csv.DictReader(open(KEY, encoding="utf-8")))
acct_strings = {r["commonwealth_account"].strip() for r in key} | {r["bd_account_id"].strip() for r in key}
acct_strings.discard("")
pk = list(csv.DictReader(open(PACKET, encoding="utf-8")))
to_confirm = {r["link_id"] for r in pk if r.get("r4_tier", "") != "DETERMINISTIC"}
all_ids = {r["link_id"] for r in pk}

if offgit:
    rest = [a for a in args if not a.lower().endswith(".xlsx")]
    inside = False
    try:
        Path(wb_path).resolve().relative_to(REPO)
        inside = True
    except ValueError:
        inside = False
    tracked = subprocess.run(["git", "ls-files", "*Match_Review*.xlsx", "*match_review*.xlsx"], capture_output=True, text=True, cwd=str(REPO)).stdout.split()
    leak = 0
    for rp in rest:
        txt = Path(rp).read_text(encoding="utf-8", errors="replace")
        leak += sum(1 for s in acct_strings if len(s) >= 6 and s in txt)
    print("SHEET-OFFGIT outside_repo=%s tracked_review_xlsx=%d record_account_strings=%d" % (not inside, len(tracked), leak))
    raise SystemExit(0)

wb = openpyxl.load_workbook(wb_path)
names = wb.sheetnames
read_me = wb[names[0]]
text_len = sum(len(str(c.value)) for row in read_me.iter_rows() for c in row if c.value)
sheets_ok = names[0] == "Read me" and text_len >= 600 and "To confirm" in names
ws = wb["To confirm"]
header = [str(c.value or "").strip() for c in ws[1]]
last = len(header)
hidden_last = ws.column_dimensions[openpyxl.utils.get_column_letter(last)].hidden is True
ids = [str(ws.cell(row=r, column=last).value or "").strip() for r in range(2, ws.max_row + 1)]
ids = [i for i in ids if i]
ids_match = sorted(ids) == sorted(to_confirm) and len(set(ids)) == len(ids)
ans_col = header.index("Your answer") + 1 if "Your answer" in header else 0
answer_blank = bool(ans_col) and all(ws.cell(row=r, column=ans_col).value in (None, "") for r in range(2, ws.max_row + 1))
answer_validation = bool(ans_col) and any(
    dv.type == "list" and "Yes" in str(dv.formula1) and "No" in str(dv.formula1) and "Unsure" in str(dv.formula1)
    and ws.cell(row=2, column=ans_col).coordinate in dv.sqref for dv in ws.data_validations.dataValidation)
vis_acct = vis_ids = 0
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=last - 1):
    for c in row:
        v = str(c.value).strip() if c.value is not None else ""
        if v in acct_strings:
            vis_acct += 1
        if v in all_ids:
            vis_ids += 1


def blanks(label_part):
    idx = [i + 1 for i, h in enumerate(header) if label_part in h.lower()]
    if not idx:
        return -1
    return sum(1 for r in range(2, ws.max_row + 1) if all(ws.cell(row=r, column=i).value in (None, "") for i in idx))


s_idx = header.index("Match strength") + 1 if "Match strength" in header else 0
strengths = Counter(str(ws.cell(row=r, column=s_idx).value or "").strip().lower() for r in range(2, ws.max_row + 1)) if s_idx else Counter()
print("REVIEW-SHEET sheets_ok=%s rows=%d ids_match=%s link_id_hidden_last=%s answer_blank=%s answer_validation=%s visible_account_strings=%d visible_link_ids=%d "
      "blank_old_name=%d blank_new_name=%d blank_value=%d strengths=%s"
      % (sheets_ok, len(ids), ids_match, hidden_last, answer_blank, answer_validation, vis_acct, vis_ids,
         blanks("old account name"), blanks("proposed account name"), blanks("old account value"),
         ",".join("%s:%d" % kv for kv in sorted(strengths.items()))))
