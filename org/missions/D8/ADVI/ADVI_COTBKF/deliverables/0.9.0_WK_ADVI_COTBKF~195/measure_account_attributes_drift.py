"""Measure the drift between one stale silver account_attributes partition and one bronze pull.

Worker-written measurement (round 105 CB-db-1). Reads two local copies of read-only S3 objects.
Prints ONE JSON object of COUNTS ONLY: no account id, name or dollar value is printed.

Usage:
  python -I measure_account_attributes_drift.py <silver_partition.parquet> <bronze_pull.xlsx>

Comparison rule: each value is normalized the way the silver writer stores it (None -> '',
integral float -> integer string, otherwise stripped string). Accounts are matched on the
stripped AccountID. For an account present on both sides, a column has changed when the SET of
its values differs between the two sides (a multi-book account compares its set of values).
multibook_flag_disagree_new counts bronze accounts with more than one row whose supervised,
billable or position_only values are not all equal.
"""
import json
import sys
from collections import defaultdict

import openpyxl
import pyarrow.parquet as pq

SILVER_COLS = {
    'group': 'portfolio_group',
    'acct': 'account_id',
    'pos': 'position_only',
    'sup': 'account_supervised',
    'bill': 'account_billable',
}
BRONZE_COLS = {
    'group': 'Portfolio Group',
    'acct': 'AccountID',
    'pos': 'Position Only',
    'sup': 'Account Supervised',
    'bill': 'Account Billable',
}
COMPARE_COLS = ('sup', 'bill', 'pos', 'group')
FLAG_COLS = ('sup', 'bill', 'pos')


def norm(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_silver(path):
    t = pq.read_table(path).to_pydict()
    n = len(t[SILVER_COLS['acct']])
    rows = []
    for i in range(n):
        rows.append({k: norm(t[c][i]) for k, c in SILVER_COLS.items()})
    return rows


def load_bronze(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active
    it = ws.iter_rows(values_only=True)
    header = [norm(h) for h in next(it)]
    idx = {k: header.index(c) for k, c in BRONZE_COLS.items()}
    rows = []
    for r in it:
        rows.append({k: norm(r[idx[k]]) for k in BRONZE_COLS})
    wb.close()
    return rows


def by_acct(rows):
    d = defaultdict(list)
    skipped = 0
    for r in rows:
        if r['acct'] == '':
            skipped += 1
            continue
        d[r['acct']].append(r)
    return d, skipped


def main():
    silver_path, bronze_path = sys.argv[1], sys.argv[2]
    s_rows = load_silver(silver_path)
    b_rows = load_bronze(bronze_path)
    old, old_blank = by_acct(s_rows)
    new, new_blank = by_acct(b_rows)

    old_keys = set(old)
    new_keys = set(new)
    both_keys = old_keys & new_keys

    changed = {c: 0 for c in COMPARE_COLS}
    changed_any = 0
    for a in both_keys:
        any_diff = False
        for c in COMPARE_COLS:
            so = {r[c] for r in old[a]}
            sn = {r[c] for r in new[a]}
            if so != sn:
                changed[c] += 1
                any_diff = True
        if any_diff:
            changed_any += 1

    multibook_flag_disagree_new = 0
    for a, rs in new.items():
        if len(rs) > 1:
            if any(len({r[c] for r in rs}) > 1 for c in FLAG_COLS):
                multibook_flag_disagree_new += 1

    drift = {
        'old_rows': len(s_rows) - old_blank,
        'new_rows': len(b_rows) - new_blank,
        'old_accounts': len(old_keys),
        'new_accounts': len(new_keys),
        'added': len(new_keys - old_keys),
        'removed': len(old_keys - new_keys),
        'both': len(both_keys),
        'multibook_old': sum(1 for a in old if len(old[a]) > 1),
        'multibook_new': sum(1 for a in new if len(new[a]) > 1),
        'changed_supervised': changed['sup'],
        'changed_billable': changed['bill'],
        'changed_position_only': changed['pos'],
        'changed_portfolio_group': changed['group'],
        'changed_any': changed_any,
        'multibook_flag_disagree_new': multibook_flag_disagree_new,
    }
    aux = {
        'silver_blank_account_rows_skipped': old_blank,
        'bronze_blank_account_rows_skipped': new_blank,
        'bronze_rows_raw': len(b_rows),
        'silver_rows_raw': len(s_rows),
    }
    print(json.dumps({'drift': drift, 'aux': aux}, indent=2))


if __name__ == '__main__':
    main()
