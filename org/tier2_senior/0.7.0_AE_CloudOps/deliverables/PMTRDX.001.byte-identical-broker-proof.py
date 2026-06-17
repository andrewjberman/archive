"""PMTRDX.001 binding empirical gate (CTO 152237Z) -- broker Vident Excel byte-identical proof.

RESULT (run 2026-06-17, pre-deploy): PASS -- 47 cells each, 0 diffs (raw value + number_format).
  col-6 (proposed) raw stays 5.05 / "5.05"  (the 100x risk -- NOT 0.0505) -- empirically eliminated.
  col-5 raw 0.0555 / 5.55% ; col-7 raw -0.005 / 0% ; multi-security (AMZN sell + MSFT buy) byte-identical.
  Vident CSV byte-identical by construction (generate_trade_csv reads the raw request body, unchanged).

Reproduces: patched code (working tree / >= c48ca8423) + a FRACTION record vs the PRE-PATCH code
(c48ca8423^) + the equivalent PERCENT record. Vident ingests the RAW cell value, so any divergence
(esp col-6 5.05 -> 0.0505) = a 100x broker error = REDLINE. Run from the workspace root.
"""
import importlib.util
import os
import subprocess
import sys

from openpyxl import load_workbook

os.environ.setdefault("AWS_DEFAULT_REGION", "us-east-1")
os.environ.setdefault("AWS_REGION", "us-east-1")
SRC = "infra/aws/lambda/trade-pipeline/trade-workflow-handler/src"
sys.path.insert(0, SRC)
WF = SRC + "/excel_generator.py"
PRE_PATCH = "c48ca8423^"  # parent of the PMTRDX.001 fix commit = today's (pre-fix) behavior


def load_mod(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


orig_src = subprocess.check_output(["git", "show", f"{PRE_PATCH}:{WF}"], text=True)
os.makedirs(".runtmp", exist_ok=True)
with open(".runtmp/_orig_excel_generator.py", "w", encoding="utf-8") as f:
    f.write(orig_src)
orig = load_mod("orig_eg", ".runtmp/_orig_excel_generator.py")
patched = load_mod("patched_eg", WF)


def securities(scale):
    return [
        {"ticker": "AMZN", "sedol": "2000019", "isin": "US0231351067", "figi": "",
         "current_weight": str(5.55 / scale), "proposed_weight": str(5.05 / scale),
         "delta": str(-0.5 / scale), "trade_instruction": "SELL", "comments": ""},
        {"ticker": "MSFT", "sedol": "2588173", "isin": "US5949181045", "figi": "",
         "current_weight": str(3.20 / scale), "proposed_weight": str(4.10 / scale),
         "delta": str(0.90 / scale), "trade_instruction": "BUY", "comments": ""},
    ]


def submission(scale):
    return {"fund": "Test Fund", "weights_as_of": "2026-06-17",
            "proposed_trades_as_of": "2026-06-17", "trade_comments": "", "other_comments": "",
            "securities": securities(scale)}


today_buf = orig.generate_vident_excel(submission(1))       # percent record, pre-patch code
fixed_buf = patched.generate_vident_excel(submission(100))  # fraction record, patched code


def cells(buf):
    buf.seek(0)
    ws = load_workbook(buf).active
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None or c.number_format != "General":
                out[c.coordinate] = (repr(c.value), c.number_format)
    return out


a, b = cells(today_buf), cells(fixed_buf)
keys = sorted(set(a) | set(b))
diffs = [(k, a.get(k), b.get(k)) for k in keys if a.get(k) != b.get(k)]
print(f"today(pre-patch+percent) cells={len(a)}  fixed(patched+fraction) cells={len(b)}  diffs={len(diffs)}")
for coord in ("E8", "F8", "G8", "E9", "F9", "G9"):
    print(f"  {coord}: today={a.get(coord)}  fixed={b.get(coord)}")
if diffs:
    print("REDLINE -- cell divergences:")
    for k, av, bv in diffs[:30]:
        print(f"  {k}: today={av}  fixed={bv}")
    sys.exit(1)
print("PASS: broker Vident Excel BYTE-IDENTICAL (every cell: raw value + number_format). No 100x risk.")
